"""P1 接口测试：/api/search/stats、/api/compare、/api/export。

口径（接口字典 6.1）：
- stats 只统计当前查询和当前角色可见记录，不泄露隐藏分组
- compare 最多 4 条，继续经过 RolePolicy；越权/不存在 → 404
- export XLSX 固定四个 Sheet；CSV 仅含当前角色可见的扁平公开字段
"""

from __future__ import annotations

import io
import json
import os
import sys
import unittest
from pathlib import Path
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from app.repository.sqlite_repository import SQLiteRepository
from app.search.llm_provider import NullProvider
from app.web.app import create_app
from app.web.seed import seed
from app.domain import role_policy


class P1ApiTests(unittest.TestCase):
    def setUp(self):
        self.repo = SQLiteRepository(":memory:")
        seed(self.repo)
        self.app = create_app(repository=self.repo, llm_provider=NullProvider())
        self.app.config["TESTING"] = True
        self.client = self.app.test_client()

    def tearDown(self):
        self.repo.close()
        os.environ.pop(role_policy.DEV_ROLE_SWITCH_ENABLED, None)

    def _switch(self, role):
        os.environ[role_policy.DEV_ROLE_SWITCH_ENABLED] = "true"
        resp = self.client.post("/api/dev/role", json={"role": role})
        self.assertEqual(resp.status_code, 200)

    # ---- stats ----

    def test_stats_guest_only_published(self):
        resp = self.client.get("/api/search/stats?city=成都")
        data = resp.get_json()
        self.assertEqual(resp.status_code, 200)
        self.assertLess(data["total"], 4)  # 只统计成都的 published 记录
        self.assertNotIn("review_status_distribution", data)  # 非 admin 不泄露复核分布
        edu = data["distributions"]["education"]
        self.assertEqual(sum(edu.values()), data["total"])

    def test_stats_admin_includes_review_distribution(self):
        self._switch("admin")
        data = self.client.get("/api/search/stats").get_json()
        self.assertEqual(data["total"], 7)
        self.assertIn("review_status_distribution", data)
        self.assertGreater(
            data["review_status_distribution"].get("review_required", 0), 0
        )

    def test_stats_respects_current_query(self):
        all_stats = self.client.get("/api/search/stats").get_json()
        filtered = self.client.get("/api/search/stats?major=计算机").get_json()
        self.assertLess(filtered["total"], all_stats["total"])

    # ---- compare ----

    def test_compare_ok_with_role_policy(self):
        resp = self.client.post("/api/compare", json={
            "record_keys": ["portal-10001-01", "portal-10001-02"],
        })
        data = resp.get_json()
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(data["records"]), 2)
        for dto in data["records"]:
            self.assertIn("fields", dto)
            self.assertIn("evidence", dto)
            self.assertNotIn("person", dto)  # 游客对比不泄露受控字段

    def test_compare_limit_4(self):
        keys = ["portal-10001-01"] * 5
        resp = self.client.post("/api/compare", json={"record_keys": keys})
        self.assertEqual(resp.status_code, 400)

    def test_compare_unknown_record_404(self):
        resp = self.client.post("/api/compare", json={"record_keys": ["no-such"]})
        self.assertEqual(resp.status_code, 404)

    def test_compare_draft_hidden_from_guest(self):
        resp = self.client.post("/api/compare", json={"record_keys": ["portal-10005-01"]})
        self.assertEqual(resp.status_code, 404)  # 越权即 4xx，不暴露存在性
        self._switch("admin")
        resp = self.client.post("/api/compare", json={"record_keys": ["portal-10005-01"]})
        self.assertEqual(resp.status_code, 200)

    # ---- export ----

    def test_export_csv_guest(self):
        resp = self.client.get("/api/export?format=csv")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/csv", resp.mimetype)
        self.assertIn("attachment", resp.headers["Content-Disposition"])
        text = resp.get_data(as_text=True)
        lines = text.lstrip("\ufeff").strip().splitlines()
        self.assertEqual(lines[0].split(",")[0], "记录编号")
        self.assertEqual(len(lines) - 1, 4)  # 4 条 published
        self.assertNotIn("来源链接", lines[0])  # 游客不导出需要登录的 URL
        self.assertNotIn("person_name", text)

    def test_export_csv_admin_includes_source_url(self):
        self._switch("admin")
        text = self.client.get("/api/export?format=csv").get_data(as_text=True)
        self.assertIn("来源链接", text.lstrip("\ufeff").splitlines()[0])

    def test_export_xlsx_four_sheets(self):
        resp = self.client.get("/api/export?format=xlsx")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp.mimetype)
        with ZipFile(io.BytesIO(resp.get_data())) as zf:
            from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(resp.get_data()), read_only=True)
        try:
            self.assertEqual(wb.sheetnames,
                             ["搜索结果", "字段证据", "来源文章", "导出说明"])
            rows = list(wb["搜索结果"].iter_rows(values_only=True))
            self.assertEqual(rows[0][0], "记录编号")
            self.assertEqual(len(rows) - 1, 4)
            notes = list(wb["导出说明"].iter_rows(values_only=True))
            flat = json.dumps(notes, ensure_ascii=False, default=str)
            self.assertIn("脱敏声明", flat)
        finally:
            wb.close()

    def test_export_bad_format_400(self):
        resp = self.client.get("/api/export?format=pdf")
        self.assertEqual(resp.status_code, 400)

    def test_export_unknown_param_400(self):
        resp = self.client.get("/api/export?format=csv&foo=1")
        self.assertEqual(resp.status_code, 400)


if __name__ == "__main__":
    unittest.main()
