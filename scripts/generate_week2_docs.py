from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / "docs" / "week2"


def set_cell_shading(cell, fill: str = "D9EAF7") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def base(title: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(20)
    p2 = doc.add_paragraph("第二周执行周期：2026年9月21日—9月27日 23:59；9月26日18:00进入P0冻结")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def para(doc, text):
    doc.add_paragraph(text)


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, v in enumerate(headers):
        t.rows[0].cells[i].text = v
        set_cell_shading(t.rows[0].cells[i])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return t


def save(doc, name):
    path = OUT / name
    doc.save(path)
    return path


def overall():
    d = base("第二周总体目标与协作执行手册")
    h(d, "一、总体目标")
    para(d, "本周采用“第一周缺口清零+第二周功能并行”。交付重点是验收闭环、B 定向接入、管理员发布链路、持久化用户功能、LAN/备份证据和最终演示。公网 HTTPS/CAS 只有在真实服务器、域名和回调配置齐备时推进，否则保留 LAN 证据并记录外部阻塞。")
    table(d, ["日期", "里程碑", "必须留下的证据"], [
        ("9/21", "C 发布集成基线；A LAN；B 白名单迁移", "分支 SHA、设备信息、允许目录清单"),
        ("9/22", "B 固定样本导入；C schema 草案", "导入日志、行数、重复幂等结果"),
        ("9/23", "用户表/Repository/API 第一版；A 备份恢复", "迁移 SQL、API 测试、恢复前后行数"),
        ("9/24", "个人中心；B 20 样本复评；A 生产配置", "移动截图、归一化报告、条件式部署记录"),
        ("9/25", "完整演示链路与跨角色回归", "录屏、Network 脱敏、跨用户隔离输出"),
        ("9/26 18:00", "P0 冻结", "只允许修阻断、安全、权限、移动端、演示问题"),
        ("9/27", "独立审查、合并、打标签", "审查报告、发布说明、最终录屏"),
    ])
    h(d, "二、分支与交接")
    bullets(d, ["集成基线为 feat/week1-core-search@473f812，建立 feat/week2-A-platform、feat/week2-B-quality、feat/week2-C-product、feat/week2-integration。", "B 严格只迁移 app/extraction、tests/extraction、data/gold、reports/extraction、progress/week1/B、evidence/week1/B，不能整体合并旧 B 分支。", "生产者必须提交 Commit、运行命令、实际值、证据路径；接收者记录接收 Commit、Schema/测试结果、问题和接收时间，每次交接至少保留四小时。"])
    h(d, "三、总体验收门槛")
    bullets(d, ["传统搜索、详情、对比、统计、CSV/XLSX 导出在关闭 LLM 时仍可用；游客只看 published 与派生脱敏字段。", "保存搜索只保存白名单 QueryPlan；提醒仅 off/daily/weekly；历史和推荐默认关闭，启用后最多保留90天。", "X-Incognito-Mode: 1 不写历史、筛选、点击、对比、导出或推荐信号；删除历史后不继续参与画像。", "B 本周以归一化口径作为验收主口径；同一20样本、同一 SHA-256、同一 gold 不变，严格逐样本结果保留为审计附件；多模态仍未过门槛则保持不启用。"])
    h(d, "四、最终演示")
    para(d, "依次演示 /healthz 与行数、游客传统搜索和证据、AI 搜索降级、对比导出、保存搜索/每周提醒、published 导入通知、隐私中心、无痕搜索、角色切换与越权拒绝、390×844/768×1024 页面、B 归一化评测与不启用多模态、LAN 访问和同步中断恢复。")
    return save(d, "00_第二周总体目标与协作执行手册.docx")


def role_a():
    d = base("角色A：平台、认证、部署与灾备任务书")
    h(d, "一、本周结果")
    para(d, "把系统从“本机可运行”推进到“LAN 可复测、生产配置可检查、备份可恢复”。真实 CAS/HTTPS 条件不足时必须保留阻塞记录，不伪造上线结果。")
    h(d, "二、日程与交付")
    table(d, ["日期", "工作", "交付/验收"], [
        ("9/21", "另一台手机或电脑访问 LAN 首页、/healthz、搜索、详情", "设备型号、访问时间、截图、响应状态"),
        ("9/22", "固化 AuthProvider/CAS Provider；关闭开发开关后 /api/dev/role 404", "契约文件、越权测试输出"),
        ("9/23", "SQLite 脱敏快照、备份、恢复、行数校验", "命令、SHA、articles/records/evidence 行数"),
        ("9/24", "有真实资源则推进 HTTPS/CAS，否则记录外部阻塞", "真实证据或阻塞记录"),
        ("9/25-27", "与 C 联调、录屏、问题清单", "交接确认和最终证据索引"),
    ])
    h(d, "三、验收清单")
    bullets(d, ["100篇台账、5篇固定样本、A4 7/7 场景均有可定位证据。", "另一台 LAN 设备可访问首页、传统查询、详情和 /healthz。", "开发角色开关关闭时模拟角色接口返回404；生产不接受客户端 header 伪造身份。", "备份恢复后文章、人物、证据行数一致；备份不含 Cookie、密码、API Key、真实姓名、原始海报和绝对路径。", "提交 feat(platform): add week2 auth and production guards、test(platform): verify lan backup restore and role isolation、docs(week2): record deployment and disaster recovery evidence。"])
    h(d, "四、交接格式")
    para(d, "交给 C：Commit/PR、运行命令、设备和网络信息、备份 SHA、恢复前后行数、CAS/HTTPS 是否具备。C 在验收台账填写接收 Commit、Schema 结果、测试结果和问题；缺任一项只能标记待接收。")
    return save(d, "01_角色A_平台认证部署任务书.docx")


def role_b():
    d = base("角色B：抽取质量与评测闭环任务书")
    h(d, "一、目录白名单")
    para(d, "从 C 基线建立 feat/week2-B-quality，只提交 app/extraction/**、tests/extraction/**、data/gold/**、reports/extraction/**、progress/week1/B/**、evidence/week1/B/**。根配置、C 领域/仓储/搜索/前端和文档必须人工审查。")
    h(d, "二、日程与交付")
    table(d, ["日期", "工作", "交付/验收"], [
        ("9/21", "定向迁移并列出允许目录", "白名单清单与 PR"),
        ("9/22", "固定5篇 bundle/evidence pack 导入 C SQLite 烟测", "Schema、行数、重复幂等日志"),
        ("9/23", "专业/城市/岗位通用规则修正", "规则 diff 与回归测试"),
        ("9/24", "同一20样本、SHA-256、gold 重新评测", "evaluation_report_v2.json/CSV"),
        ("9/25-26", "复核队列、授权、保留策略与游客素材检查", "review_required 清单与脱敏检查"),
        ("9/27", "演示讲解评测口径", "归一化结果与不启用多模态结论"),
    ])
    h(d, "三、统一验收口径")
    para(d, "本周验收只把归一化口径作为主判定：五字段逐项记录分子、分母、归一化准确率和低于90%的样本；缺失、冲突、低置信度进入 review_required。原始逐样本结果与严格口径文件不得删除，作为可追溯附件。不得改变20样本集合、SHA-256、gold 或多模态启用门槛。")
    h(d, "四、交接给C")
    bullets(d, ["固定5篇 bundle 和 evidence pack 可被导入器接收，重复导入不增加记录。", "给出报告、成本、模型、授权与保留策略，以及复核队列。", "游客服只能拿派生图或安全占位图，不能原图下发后再靠 CSS 遮挡。", "提交 feat(extraction): integrate bundles and evidence packs on week2 base、test(extraction): rerun fixed gold evaluation without sample changes、docs(extraction): record normalized quality gate。"])
    return save(d, "02_角色B_抽取质量与评测任务书.docx")


def role_c():
    d = base("角色C：用户功能与最终整合任务书")
    h(d, "一、本周必须落地")
    bullets(d, ["保存搜索、提醒、系统眼中的我、隐私中心、无痕模式服务端约束。", "接收 B 白名单成果并完成固定样本导入烟测；补管理员发布/撤回和审计事件。", "个人中心、保存搜索列表、提醒开关、通知列表、隐私中心和移动端页面。"])
    h(d, "二、固定接口")
    para(d, "GET/POST/DELETE /api/saved-searches；PATCH /api/saved-searches/{id}/alert；GET/PATCH /api/me/privacy；DELETE /api/me/history；GET /api/me/insights；GET /api/notifications。保存搜索只保存白名单 QueryPlan；提醒 off/daily/weekly，默认 weekly；历史/推荐默认关闭，启用后90天。")
    h(d, "三、验收清单")
    bullets(d, ["同一用户可创建、读取、修改提醒和删除保存搜索；两个用户互相不可读取保存搜索、历史、通知或画像。", "X-Incognito-Mode: 1 时无任何个人持久化写入；删除历史后旧记录不参与推荐。", "游客响应不含姓名、原图引用、联系方式、二维码、会议入口和受限来源 URL；review_required 不进入游客/普通校内查询。", "关闭 LLM 后传统搜索、详情、对比、统计、导出不回归；四 Sheet 导出、既有性能阈值和移动布局通过。", "9/26 18:00 冻结 P0；9/27 合并、打标签、发布说明、录屏。"])
    h(d, "四、C 的接收与最终演示")
    para(d, "每次接收记录生产 Commit、Schema 结果、测试结果、问题和确认时间。最终演示要覆盖健康检查、传统/AI 降级、对比导出、保存搜索和提醒、隐私/无痕、角色越权、移动端、B 归一化评测和 LAN 恢复。")
    return save(d, "03_角色C_用户功能与最终整合任务书.docx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for fn in (overall, role_a, role_b, role_c):
        print(fn())
