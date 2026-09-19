"""金标准 v2 构建：从用户填写的工作簿转录为 data/gold/gold_20.json。

背景（PROGRESS 2026-09-19 选样 v2 裁定）：
- 用户以标注人身份交付 20 张经验分享类海报的金标准（40 行人物标注）；
- 用户表的 P 编号与初版冻结清单不一致，本脚本按海报内容建立映射：
  14 张台账经验分享（去掉校友就业帖 246444 与 5 张公告类）+ 6 张场外补充；
- 新样本 ID 按 notice_id 确定性排序（场外用 manual-* ID 排在台账后）；
- 字段值逐字转录，不做任何改写；疑似笔误进入差异报告由标注人裁决；
- 仓库内 JSON 不含姓名（PII 只存受控目录）。

运行：python -m app.extraction.build_gold_v2 <用户工作簿.xlsx>
"""

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]

# 用户表 (样本ID, 行号) -> 海报 notice_id（内容核对见 PROGRESS 选样 v2 一节）
ROW_MAP = {
    ('P01', 1): '246592', ('P01', 2): '246592',      # 广西
    ('P02', 1): '246193', ('P02', 2): '246193',      # 山东
    ('P03', 1): '246323', ('P03', 2): '246323',      # 青海
    ('P04', 1): '246590', ('P04', 2): '246590',      # 云南
    ('P05', 1): '246674', ('P05', 2): '246674',      # 天津
    ('P06', 1): '246709', ('P06', 2): '246709',      # 河南
    ('P07', 1): '246757', ('P07', 2): '246757',      # 宁夏
    ('P08', 1): '246951', ('P08', 2): '246951',      # 湖南
    ('P09', 1): '247035', ('P09', 2): '247035',      # 福建
    ('P10', 1): '247339', ('P10', 2): '247339',      # 广东
    ('P11', 1): '247509', ('P11', 2): '247509',      # 辽宁
    ('P12', 1): '247746', ('P12', 2): '247746',      # 山西 1月2日
    ('P13', 1): '247919', ('P13', 2): '247919',      # 河北 1月13日
    ('P14', 1): 'manual-jiangxi-1024', ('P14', 2): 'manual-jiangxi-1024',
    ('P15', 1): 'manual-guizhou-1018', ('P15', 2): 'manual-guizhou-1018',
    ('P16', 1): 'manual-xinjiang-0314', ('P16', 2): 'manual-xinjiang-0314',
    ('P17', 1): 'manual-hebei-0312', ('P17', 2): 'manual-hebei-0312',
    ('P18', 1): 'manual-shanxi-1227', ('P18', 2): 'manual-shanxi-1227',
    ('P19', 1): 'manual-shanxi-1227',   # 山西12/27 第三人（用户排在 P19 首行）
    ('P19', 2): 'manual-hubei-1219',    # 湖北12/19 吴
    ('P20', 1): '247586',               # 新疆12/23（单人帖）
    ('P20', 2): 'manual-hubei-1219',    # 湖北12/19 刘
}

# 场外海报元数据（标题取自门户页面页眉；发布时间取页面"发布时间"）
MANUAL_META = {
    'manual-xinjiang-0314': ('3月14日||"共美"沙龙—新疆维吾尔自治区定向选调经验分享',
                             '2024-03-07 10:42'),
    'manual-guizhou-1018': ('10月18日||"共美"沙龙——贵州定向选调经验分享',
                            '2024-10-16 10:05'),
    'manual-jiangxi-1024': ('10月24日||"共美"沙龙——江西定向选调经验分享',
                            '2024-10-16 10:16'),
    'manual-shanxi-1227': ('12月27日||"共美"沙龙——山西定向选调经验分享',
                           '2023-12-22 19:23'),
    'manual-hubei-1219': ('12月19日||"共美"沙龙——湖北选调经验分享',
                          '2023-12-15 13:35'),
    'manual-hebei-0312': ('3月12日||"共美"沙龙—河北定向选调经验分享',
                          '2024-03-05 16:09'),
}

# 场外海报图片目录（batch2 提取时的暂存名 -> manual ID）
MANUAL_ASSET_SRC = {
    'manual-xinjiang-0314': 'x-0314-xinjiang',
    'manual-guizhou-1018': 'x-1018-guizhou',
    'manual-jiangxi-1024': 'x-1024-jiangxi',
    'manual-shanxi-1227': 'x-1227-shanxi',
    'manual-hubei-1219': 'x-1219-hubei',
    'manual-hebei-0312': 'x-0312-hebei',
}

FIELDS = ('cohort', 'grade', 'education', 'college', 'major',
          'city', 'position_or_unit')


def read_user_rows(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb['金标准标注']
    rows = []
    for row in ws.iter_rows(min_row=5):
        vals = [c.value for c in row]
        pid = vals[1]
        if pid is None:
            continue
        seq = int(vals[2] or 0)
        fields = dict(zip(FIELDS, vals[4:11]))
        rows.append({'user_pid': str(pid).strip(), 'seq': seq,
                     'fields': {k: (str(v).strip() if v is not None else None)
                                for k, v in fields.items()}})
    return rows


def main(argv: list[str]) -> int:
    wb_path = Path(argv[0])
    base = ROOT.parent
    controlled = base / 'controlled_assets'

    # 1) 读用户标注并按海报归组
    user_rows = read_user_rows(wb_path)
    assert len(user_rows) == 40, f'期望 40 行人物标注，实际 {len(user_rows)}'
    by_notice: dict[str, list[dict]] = {}
    for r in user_rows:
        key = (r['user_pid'], r['seq'])
        if key not in ROW_MAP:
            raise SystemExit(f'未映射的标注行: {key}')
        by_notice.setdefault(ROW_MAP[key], []).append(r['fields'])

    # 2) 新样本 ID：台账 14 张按 notice_id 排序，场外 6 张按 manual ID 排序
    ledger_ids = sorted(n for n in by_notice if not n.startswith('manual-'))
    manual_ids = sorted(n for n in by_notice if n.startswith('manual-'))
    ordered = ledger_ids + manual_ids
    assert len(ordered) == 20, f'海报数应为 20，实际 {len(ordered)}'

    # 3) 哈希：台账取自 selection v1，场外取自 batch2 提取报告
    sel = json.loads((ROOT / 'data' / 'gold' / 'sample_selection_20.json')
                     .read_text(encoding='utf-8'))
    hash_by_nid = {s['notice_id']: s.get('image_sha256')
                   for s in sel['samples']}
    batch2 = {r['notice_id']: r['sha256'] for r in
              json.loads((controlled / 'batch2_report.json').read_text(encoding='utf-8'))
              if 'notice_id' in r}
    for manual_id, staging in MANUAL_ASSET_SRC.items():
        if staging in batch2:
            hash_by_nid.setdefault(manual_id, batch2[staging])

    # 4) 组装 gold（逐字转录，不含姓名）
    samples = []
    for i, nid in enumerate(ordered, start=1):
        persons = by_notice[nid]
        sha = hash_by_nid.get(nid)
        if not sha:
            raise SystemExit(f'{nid} 缺 image_sha256')
        samples.append({
            'sample_id': f'P{i:02d}',
            'notice_id': nid,
            'image_sha256': sha,
            'annotated_by': 'annotator-B',
            'evidence_summary': '海报人物信息块（姓名徽章下方文字行）',
            'fields': persons[0],          # 兼容单人/首人物；多人见 persons
            'persons': [dict(p) for p in persons],
            'is_multi_person': len(persons) > 1,
        })

    gold = {
        'sample_set': 'week1-poster-gold-20-v2',
        'annotator': 'annotator-B（人工盲填，工作簿原表存档于受控目录）',
        'frozen_at': None,   # 差异裁决完成后回填
        'revision': 'v2：20 张全部为经验分享类人物海报。初版清单中的校友就业帖'
                    '(246444) 与 5 张公告类移出（无人物五字段），6 张台账外经验'
                    '分享补充入场（manual-*，来源受控交接 PDF，手工建 bundle）。',
        'notes': '字段值逐字转录自标注人工作簿；用户表 P 编号与本清单不一致，'
                 '映射关系见 app/extraction/build_gold_v2.py ROW_MAP（按海报内容核对）。',
        'samples': samples,
    }

    # 5) 机器校验（复用校验器）+ 写文件
    from data.gold.gold_schema_check import validate_gold
    errors = validate_gold(gold)
    if errors:
        print('GOLD 校验失败:')
        for e in errors:
            print(' -', e)
        return 1

    out = ROOT / 'data' / 'gold' / 'gold_20.json'
    out.write_text(json.dumps(gold, ensure_ascii=False, indent=2), encoding='utf-8')

    # 6) selection v2（v1 备份留档）
    v1 = ROOT / 'data' / 'gold' / 'sample_selection_20_v1.json'
    if not v1.exists():
        shutil.copy(ROOT / 'data' / 'gold' / 'sample_selection_20.json', v1)
    sel_v2 = {
        'sample_set': 'week1-poster-gold-20-v2',
        'frozen': True,
        'frozen_basis': '选样 v2 由标注人裁定（见 gold revision）；'
                        'ID 按 notice_id 确定性排序',
        'selection_rule': '14 张台账经验分享（246193,246323,246590,246592,246674,'
                          '246709,246757,246951,247035,247339,247509,247586,247746,'
                          '247919）+ 6 张场外 manual-*；排除校友就业帖与公告类',
        'samples': [{'sample_id': s['sample_id'], 'notice_id': s['notice_id'],
                     'image_sha256': s['image_sha256'],
                     'is_multi_person': s['is_multi_person']}
                    for s in samples],
    }
    (ROOT / 'data' / 'gold' / 'sample_selection_20.json').write_text(
        json.dumps(sel_v2, ensure_ascii=False, indent=2), encoding='utf-8')

    # 7) 20 张评测输入 bundle（14 台账 + 6 手工）+ 场外图片目录对齐
    inputs_dir = ROOT / 'reports' / 'extraction' / 'eval_inputs'
    inputs_dir.mkdir(parents=True, exist_ok=True)
    ledger = json.loads((ROOT / 'evidence' / 'week1' / 'A' / 'ledger'
                         / 'article_ledger.json').read_text(encoding='utf-8'))
    ledger_by_nid = {e['notice_id']: e for e in ledger['entries']}
    posters_dir = controlled / 'posters'
    now = '2026-09-19T18:00:00+08:00'
    for s in samples:
        nid = s['notice_id']
        if nid.startswith('manual-'):
            title, published = MANUAL_META[nid]
            src_url = f'manual://{nid}'
        else:
            e = ledger_by_nid[nid]
            title = e['title']
            published = e['first_seen_at']  # 台账无发布时间，取采集时间（仅评测用）
            src_url = f'https://my.muc.edu.cn/notice/{nid}'
        bundle = {
            'schema_version': 'article_bundle.v1', 'notice_id': nid,
            'title': title, 'source_url': src_url,
            'published_at': published, 'content_type': 'poster',
            'clean_text': None,
            'asset_refs': [{'asset_id': f'{nid}-poster-01', 'kind': 'poster',
                            'local_ref': f'private://{nid}/poster-01',
                            'sha256': s['image_sha256']}],
            'fetch_status': 'processed', 'failure_reason': None,
            'fetched_at': now,
        }
        (inputs_dir / f'{nid}.json').write_text(
            json.dumps(bundle, ensure_ascii=False, indent=2), encoding='utf-8')
        # 场外图片目录对齐 manual-*
        if nid.startswith('manual-'):
            src_dir = posters_dir / MANUAL_ASSET_SRC[nid]
            dst_dir = posters_dir / nid
            dst_dir.mkdir(parents=True, exist_ok=True)
            for img in src_dir.glob('poster-01.*'):
                shutil.copy(img, dst_dir / img.name)

    print(f'gold_20.json 写入完成：{len(samples)} 样本 / '
          f'{sum(len(s["persons"]) for s in samples)} 人物，机器校验通过')
    for s in samples:
        print(f"  {s['sample_id']} {s['notice_id']:<22} "
              f"{'多人x%d' % len(s['persons']) if s['is_multi_person'] else '单人'}")
    return 0


if __name__ == '__main__':
    raise SystemExit(main(sys.argv[1:]))
